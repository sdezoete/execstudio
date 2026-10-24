"""End-to-end API tests with a mock LLM. Standard library only.

    python3 -m unittest discover -s tests -v
"""
import io
import json
import re
import os
import sys
import threading
import unittest
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import mock_llm  # noqa: E402

BRIEF = {"task": "new", "topic": "Q3 business performance review", "decision": "Approve a 3% mid-market price increase",
         "facts": "Revenue +4% vs plan. Gross margin -1.8 pts.", "audience": "executive", "length": "mid",
         "source": "docs", "template": "corporate"}


def _start(server):
    t = threading.Thread(target=server.serve_forever, daemon=True)
    t.start()
    return server


class Base(unittest.TestCase):
    llm_up = True

    @classmethod
    def setUpClass(cls):
        from http.server import ThreadingHTTPServer
        cls.mock = _start(mock_llm.serve(0)) if cls.llm_up else None
        llm_url = "http://127.0.0.1:%d/v1" % cls.mock.server_address[1] if cls.mock else "http://127.0.0.1:9/v1"
        os.environ["STUDIO_LLM_BASE_URL"] = llm_url
        os.environ["STUDIO_LLM_TIMEOUT"] = "5"
        os.environ["STUDIO_AUDIT_LOG"] = str(ROOT / "logs" / "test-audit.log")
        os.environ["STUDIO_SERVER_LOG"] = ""
        os.environ["STUDIO_WORK_DIR"] = str(ROOT / "work" / "test")
        import server as srv
        from studio import config
        from studio.services import Studio
        srv.Handler.studio = Studio(config.load())
        cls.app = _start(ThreadingHTTPServer(("127.0.0.1", 0), srv.Handler))
        cls.base = "http://127.0.0.1:%d" % cls.app.server_address[1]

    @classmethod
    def tearDownClass(cls):
        cls.app.shutdown()
        cls.app.server_close()
        if cls.mock:
            cls.mock.shutdown()
            cls.mock.server_close()

    def get(self, path):
        with urllib.request.urlopen(self.base + path, timeout=10) as r:
            return r.status, r.headers, r.read()

    def post(self, path, obj=None, raw=None, headers=None):
        data = raw if raw is not None else json.dumps(obj or {}).encode()
        req = urllib.request.Request(self.base + path, data=data, method="POST",
                                     headers=headers or {"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req, timeout=20) as r:
                return r.status, r.headers, r.read()
        except urllib.error.HTTPError as e:
            return e.code, e.headers, e.read()


class WithLLM(Base):
    llm_up = True

    def test_static_and_security_headers(self):
        st, h, body = self.get("/")
        self.assertEqual(st, 200)
        self.assertIn(b"Executive Presentation Studio", body)
        self.assertIn("default-src 'self'", h["Content-Security-Policy"])
        for p in ("/static/js/app.js", "/static/js/api.js", "/static/js/thumbs.js", "/static/css/app.css", "/static/img/icon.svg"):
            self.assertEqual(self.get(p)[0], 200, p)

    def test_wrong_method_answers_json(self):  # round 3 carry-over: was an HTML 501 page
        for method in ("PUT", "DELETE", "PATCH"):
            req = urllib.request.Request(self.base + "/api/storyline", data=b"{}", method=method)
            with self.assertRaises(urllib.error.HTTPError) as cm:
                urllib.request.urlopen(req, timeout=10)
            self.assertEqual(cm.exception.code, 405)
            self.assertIn("application/json", cm.exception.headers["Content-Type"])
            self.assertIn("GET, HEAD, POST", cm.exception.headers["Allow"])
            self.assertIn("not allowed", json.loads(cm.exception.read())["error"])

    def test_path_traversal_blocked(self):
        try:
            st = self.get("/static/../server.py")[0]
        except urllib.error.HTTPError as e:
            st = e.code
        self.assertEqual(st, 404)

    def test_meta_health(self):
        meta = json.loads(self.get("/api/meta")[2])
        self.assertTrue(meta["templates"] and meta["audiences"])
        health = json.loads(self.get("/api/health")[2])
        self.assertTrue(health["llm"]["reachable"])
        self.assertIn("engine", health["pptx"])

    def test_storyline_review_qa(self):
        st, _, body = self.post("/api/storyline", {"brief": BRIEF, "version": "full", "context": ""})
        self.assertEqual(st, 200, body)
        s = json.loads(body)
        self.assertEqual(s["source"], "llm")
        self.assertGreaterEqual(len(s["slides"]), 3)
        st, _, body = self.post("/api/review", {"brief": BRIEF, "slides": s["slides"]})
        r = json.loads(body)
        self.assertTrue(any(not c["ok"] for c in r["checks"]))  # slide 3 has a placeholder + missing evidence
        st, _, body = self.post("/api/qa", {"brief": BRIEF, "slides": s["slides"]})
        self.assertEqual(json.loads(body)["source"], "llm")

    def test_no_facts_gives_skeleton_not_refusal(self):  # D-09
        brief = dict(BRIEF, facts="", decision="", source="scratch")
        s = json.loads(self.post("/api/storyline", {"brief": brief, "version": "full"})[2])
        self.assertGreaterEqual(len(s["slides"]), 5)                      # a real skeleton, not one "send facts" slide
        for sl in s["slides"]:
            self.assertNotIn("I cannot", sl["notes"])
            self.assertNotRegex(sl["evidence"].get("source", ""), "(?i)source documents")
        from studio.checks import clean_storyline
        out = clean_storyline([{"notes": "Please share the data.", "evidence": {"status": "linked", "source": "SOURCE DOCUMENTS"}}], False)
        self.assertEqual((out[0]["notes"], out[0]["evidence"]["status"]), ("", "needed"))
        out = clean_storyline([{"notes": "", "evidence": {"status": "needed", "source": "No source documents provided", "gap": "No facts"}}], False)
        self.assertEqual(out[0]["evidence"]["source"], "")                  # seen live in round 7

    def test_storyline_long_headlines_shortened(self):  # retest side note: a 16-word generated headline
        topic = "quarterly review of warehouse safety incidents and the corrective actions taken across all regional sites"
        mock_llm.SHARPEN_REPLIES[:] = ["Warehouse safety incidents fell after corrective actions", "Next: roll corrective actions out to every site"]
        s = json.loads(self.post("/api/storyline", {"brief": dict(BRIEF, topic=topic, facts="Incidents down 12%. Audits up 8%."), "version": "full"})[2])
        mock_llm.SHARPEN_REPLIES[:] = []
        self.assertEqual(s["source"], "llm")
        long_ones = [x["headline"] for x in s["slides"] if x["kind"] != "title" and len(x["headline"].split()) > 15]
        self.assertEqual(long_ones, [])

    def test_review_flags_phantom_gap(self):  # D-06
        slide = {"kind": "driver", "headline": "Margin fell 2.1 pts on freight and raw materials",
                 "body": ["Freight +0.9 pts", "Raw materials +1.2 pts", "Remaining gap not specified"],
                 "evidence": {"status": "linked", "source": "Finance"}}
        r = json.loads(self.post("/api/review", {"brief": BRIEF, "slides": [slide]})[2])
        check = next(c for c in r["checks"] if c["title"] == "Figures add up")
        self.assertFalse(check["ok"])
        self.assertIn("0.9 + 1.2 = 2.1 pts", check["detail"])
        slide["body"] = slide["body"][:2]
        r = json.loads(self.post("/api/review", {"brief": BRIEF, "slides": [slide]})[2])
        self.assertTrue(next(c for c in r["checks"] if c["title"] == "Figures add up")["ok"])
        # N-02: actual / plan / variance is not a bridge, on any kind of slide
        for kind, head in (("bars", "Revenue missed plan by EUR 2.8m"), ("driver", "Revenue is EUR 48.2m, EUR 2.8m short of plan")):
            vs = {"kind": kind, "headline": head, "evidence": {"status": "linked", "source": "Finance"},
                  "body": ["Revenue EUR 48.2m vs plan EUR 51.0m", "Gap of EUR 2.8m (-5.5%)", "Plan was EUR 51m"]}
            r = json.loads(self.post("/api/review", {"brief": BRIEF, "slides": [vs]})[2])
            self.assertTrue(next(c for c in r["checks"] if c["title"] == "Figures add up")["ok"], (kind, r["checks"]))

    def test_figures_add_up_both_directions(self):  # R-01: catches real mismatches, no false alarms (N-02)
        from studio.review import _sum_problem
        s = lambda kind, head, body, notes="": {"kind": kind, "headline": head, "body": body, "notes": notes}  # noqa: E731
        must_fail = [
            s("driver", "Freight and raw materials explain the full 2.1-point gross margin decline",
              ["Freight contributed 0.9 points", "Raw materials contributed 1.6 points", "0.9 plus 1.6 equals the 2.1-point decline"]),
            s("bars", "Total revenue reached EUR 48.2m in Q3", ["North EUR 20m", "South EUR 15m", "West EUR 10m"]),
            s("summary", "Revenue is spread across three regions", ["North EUR 20m", "South EUR 15m", "West EUR 10m", "Total EUR 48.2m"]),
            s("summary", "Enterprise leads the revenue mix", ["Enterprise 55% of revenue", "Mid-market 35%", "SMB 20%"]),
            s("driver", "Margin fell 1.8 pts", ["Freight 0.9 pts", "Materials 1.2 pts"]),
            # round 4: point and currency bridges, in every spelling the tester tried, on any slide kind
            s("summary", "Freight and raw materials drove the decline", ["Gross margin down 2.1 pts", "Freight +0.9 pts", "Raw materials +1.6 pts"]),
            s("summary", "Freight and raw materials drove the decline", ["Gross margin down 2.1 pts", "Freight -0.9pp", "Raw materials -1.6pp"]),
            s("summary", "Freight and raw materials drove the decline", ["Gross margin down 2.1 points", "Freight: 0.9 points", "Raw materials: 1.6 points"]),
            s("summary", "Margin decline", ["Gross margin fell 2.1 points", "Freight contributed 0.9 points", "Raw materials contributed 1.6 points"]),
            s("bars", "Operating profit fell EUR 12.0m", ["Volume –EUR 5m", "Price EUR –3m", "Costs EUR –2m"]),
            # quality-check suggestions: parts that name the change they feed, and a YoY total
            s("summary", "Margin decline explained", ["Gross margin down 2.1 pts", "Freight added 0.9 points to the decline",
                                                      "Raw materials added 1.6 points to the decline"]),
            s("summary", "EBIT fell EUR 12.0m", ["Lower volume reduced EBIT by EUR 5.0m", "Price cut EBIT by EUR 3.0m", "Costs reduced EBIT by EUR 2.0m"]),
            s("summary", "Freight and materials explain the margin move", ["Gross margin 31.4%, -2.1 pts YoY", "Freight contributed 0.9 pts",
                                                                           "Raw materials contributed 1.6 pts"]),
            # round 5: a one-line "and" sum, and a total "vs prior year"
            s("driver", "Freight and raw materials drove the decline", ["Freight 0.9 and raw materials 1.6 points sum to the 2.1-point decline"]),
            s("summary", "Margin decline explained", ["Gross margin down 2.1 pts vs prior year", "Freight contributed 0.9 pts", "Raw materials contributed 1.6 pts"]),
            s("summary", "EBIT fell EUR 12.0m", ["Lower volume reduced EBIT by EUR 5.0m", "Price -3m", "Costs -2m"]),
        ]
        must_pass = [
            s("driver", "Freight and raw materials explain the full 2.1-point decline",
              ["Freight 0.9 points", "Raw materials 1.2 points", "0.9 plus 1.2 equals the 2.1-point decline"]),
            s("driver", "Revenue is EUR 48.2m, EUR 2.8m short of plan", ["Revenue EUR 48.2m vs plan EUR 51.0m", "Gap of EUR 2.8m (-5.5%)"]),
            s("summary", "Revenue up 4% overall", ["Enterprise +12%", "SMB -3%"]),
            s("driver", "Margin fell 2.5 pts", ["Freight 0.9 pts", "Materials 1.2 pts"], "Other factors explain the rest."),
            s("bars", "Total revenue reached EUR 45m in Q3", ["North EUR 20m", "South EUR 15m", "West EUR 10m"]),
            s("summary", "Enterprise leads the revenue mix", ["Enterprise 55% of revenue", "Mid-market 35%", "SMB 10%"]),
            s("kpi", "Revenue and margin missed plan", ["Revenue EUR 48.2m vs EUR 51.0m", "Margin 31.4%, down 2.1 pts", "Churn 4.1% vs 3.5%"]),
            s("summary", "Freight and raw materials drove the decline", ["Gross margin down 2.1 pts", "Freight +0.9 pts", "Raw materials +1.2 pts"]),
            s("bars", "Operating profit fell EUR 10.0m", ["Volume –EUR 5m", "Price EUR –3m", "Costs EUR –2m"]),
            s("summary", "Operating profit fell EUR 4m, driven by volume", ["Volume -5m", "Price +2m", "Mix -1m"]),     # mixed signs net
            s("kpi", "Margin and churn moved the wrong way", ["Gross margin down 2.1 pts", "Churn up 0.6 pts", "NPS down 3 pts"]),  # not a bridge
            # a KPI line with its own cause ("driven by", "due to") is a separate movement, not a part
            s("kpi", "Three KPIs moved", ["Gross margin down 2.1 pts due to freight", "Churn up 0.6 pts driven by the price rise", "NPS down 3 pts"]),
            s("summary", "Margin decline explained", ["Gross margin down 2.1 pts", "Freight added 0.9 points to the decline",
                                                      "Raw materials added 1.2 points to the decline"]),
            s("summary", "EBIT fell EUR 10.0m", ["Lower volume reduced EBIT by EUR 5.0m", "Price cut EBIT by EUR 3.0m", "Costs reduced EBIT by EUR 2.0m"]),
            s("kpi", "KPIs year on year", ["Margin -2.1 pts YoY", "Churn +0.6 pts YoY", "NPS -3 pts YoY"]),
            # round 5: actions on something else are not parts of the total
            s("summary", "Q3 revenue update", ["Revenue fell EUR 4.0m", "We cut costs by EUR 1.5m", "We raised prices, adding EUR 0.5m"]),
            s("summary", "Q3 revenue update", ["Revenue fell EUR 4.0m", "We cut costs by EUR 1.5m", "We raised prices by EUR 0.5m"]),
            s("summary", "Q3 revenue update", ["Revenue fell EUR 4.0m, driven by volume", "We cut costs by EUR 1.5m", "We raised prices by EUR 0.5m"]),
            s("summary", "EBIT fell EUR 10.0m", ["Lower volume reduced EBIT by EUR 5.0m", "Price cut EBIT by EUR 3.0m", "Costs reduced EBIT by EUR 2.0m"]),
            s("driver", "Freight and raw materials drove the decline", ["Freight 0.9 and raw materials 1.2 points sum to the 2.1-point decline"]),
            s("summary", "Margin decline explained", ["Gross margin down 2.1 pts vs prior year", "Freight contributed 0.9 pts", "Raw materials contributed 1.2 pts"]),
            s("driver", "Revenue down EUR 2.8m vs plan", ["Revenue EUR 48.2m vs plan EUR 51.0m", "Gap of EUR 2.8m (-5.5%)", "Plan was EUR 51m"]),
            s("summary", "Freight and raw materials drove the decline", ["Gross margin down 2.1 pts", "Freight +0.94 pts", "Raw materials +1.18 pts"]),
            # seen live: a body line restating the total is not a third part
            s("driver", "Freight and raw materials explain the full 2.1-point gross margin decline",
              ["Freight added 0.9 pts", "Raw materials added 1.2 pts", "Together they explain the full 2.1-point decline"]),
        ]
        for sl in must_fail:
            self.assertTrue(_sum_problem(sl), sl)
        for sl in must_pass:
            self.assertEqual(_sum_problem(sl), "", sl)

    def test_storyline_invented_figures(self):  # round 6: an invented plan when the brief had no facts
        for mode, calls, flagged in (("first", 2, False), ("always", 2, True)):
            mock_llm.STORY_INVENT, mock_llm.STORY_CALLS[:] = mode, []
            try:
                s = json.loads(self.post("/api/storyline", {"brief": BRIEF, "version": "full"})[2])
            finally:
                mock_llm.STORY_INVENT = ""
            self.assertEqual(len(mock_llm.STORY_CALLS), calls, mode)          # asked once, re-asked once
            last = s["slides"][-1]
            if flagged:                                                     # still invented after the retry
                self.assertEqual(last["evidence"]["status"], "needed")
                self.assertIn("15", last["evidence"]["gap"])
                self.assertIn("2028", last["evidence"]["gap"])
            else:
                self.assertNotIn("15 stores", " ".join(last["body"]))
        mock_llm.STORY_CALLS[:] = []
        self.post("/api/storyline", {"brief": BRIEF, "version": "full"})
        self.assertEqual(len(mock_llm.STORY_CALLS), 1)                      # a faithful answer is not re-asked

    def test_round6_arabic_offset_chart(self):
        from studio.review import _sum_problem
        from studio import schema
        s = lambda k, h, b: {"kind": k, "headline": h, "body": b, "notes": ""}  # noqa: E731
        # offset bridges subtract the offsetting part
        self.assertEqual(_sum_problem(s("summary", "EBIT rose SAR 12m", ["Volume added SAR 15m", "Price added SAR 6m", "Higher costs offset SAR 9m"])), "")
        self.assertTrue(_sum_problem(s("summary", "EBIT rose SAR 12m", ["Volume added SAR 15m", "Price added SAR 6m", "Higher costs offset SAR 4m"])))
        # the figures check on Arabic slides
        self.assertTrue(_sum_problem(s("summary", "انخفض هامش الربح الإجمالي ٢٫١ نقطة", ["ساهم الشحن بـ ٠٫٩ نقطة", "ساهمت المواد الخام بـ ١٫٦ نقطة"])))
        self.assertEqual(_sum_problem(s("summary", "انخفض هامش الربح ٢٫١ نقطة", ["ساهم الشحن بـ ٠٫٩ نقطة", "ساهمت المواد الخام بـ ١٫٢ نقطة"])), "")
        self.assertTrue(_sum_problem(s("bars", "إجمالي الإيرادات ٤٨٫٢ مليون ريال", ["الشمال ٢٠ مليون ريال", "الجنوب ١٥ مليون ريال", "الغرب ١٠ مليون ريال"])))
        self.assertTrue(_sum_problem(s("summary", "حصة المؤسسات هي الأكبر", ["حصة المؤسسات ٥٥٪ من الإيرادات", "الشركات المتوسطة ٣٥٪", "الشركات الصغيرة ٢٠٪"])))
        self.assertEqual(_sum_problem(s("driver", "انخفضت الإيرادات ٢٫٨ مليون ريال عن الخطة",
                                        ["الإيرادات ٤٨٫٢ مليون ريال مقابل الخطة ٥١ مليون ريال", "فجوة ٢٫٨ مليون ريال"])), "")
        self.assertEqual(_sum_problem(s("kpi", "مؤشرات الربع الثالث", ["انخفض الهامش ٢٫١ نقطة", "ارتفع معدل التسرب ٠٫٦ نقطة", "انخفض مؤشر الرضا ٣ نقاط"])), "")
        # one unit per chart
        ch = schema.chart({"categories": ["Q1", "Q2"], "series": [{"name": "Revenue (SAR m)", "values": [48.2, 51]},
                                                                  {"name": "Gross margin %", "values": [31.4, 33.5]}]})
        self.assertEqual([x["name"] for x in ch["series"]], ["Revenue (SAR m)"])

    def test_arabic_export_labels_and_filename(self):  # round 6: English labels and "____2026" names
        from pptx import Presentation
        slides = [{"kind": "summary", "headline": "انخفضت الحوادث بعد التدريب", "body": ["نقطة %d" % k for k in range(1, 8)],
                   "evidence": {"status": "linked", "source": "تقرير السلامة"}},
                  {"kind": "roadmap", "headline": "الخطوات التالية", "body": ["تقييم المواقع", "تدريب الفرق", "مراجعة النتائج"]}]
        st, h, body = self.post("/api/export/pptx", {"brief": BRIEF, "title": "مراجعة السلامة", "slides": slides})
        self.assertEqual(st, 200)
        self.assertIn("presentation_", h["Content-Disposition"].split("filename*")[0])  # readable ASCII fallback
        self.assertIn("UTF-8''%D9%85", h["Content-Disposition"])                         # the Arabic name itself
        prs = Presentation(io.BytesIO(body))
        texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
        notes = " ".join(s.notes_slide.notes_text_frame.text for s in prs.slides)
        for english in ("Source:", "Phase 1", "Moved from the slide"):
            self.assertNotIn(english, texts + notes)
        self.assertIn("المصدر:", texts)
        self.assertIn("المرحلة 1", texts)
        self.assertIn("نُقل من الشريحة:", notes)

    def test_summary_keeps_criticism_drops_praise(self):  # quality-check suggestion on S-01
        from studio.services import _praises
        self.assertFalse(_praises("Evidence is thin on slides 3 and 4."))
        self.assertFalse(_praises("Evidence is not strong enough for a board decision."))
        self.assertFalse(_praises("Evidence is thin as well."))
        self.assertTrue(_praises("Decision clarity is high, with no ambiguity about the ask."))
        self.assertTrue(_praises("Decision clarity is strong, though the timeline could be tighter."))
        self.assertTrue(_praises("Evidence is well sourced."))

    def test_advise_and_sharpen(self):
        a = json.loads(self.post("/api/advise", {"message": "why did turnover increase"})[2])
        self.assertEqual(a["primary"]["kind"], "driver")
        h = json.loads(self.post("/api/sharpen", {"headline": "Margin overview", "body": []})[2])
        self.assertIn("1.8", h["headline"])

    def test_sharpen_holds_the_15_word_rule(self):  # D-04
        long_one = "Key performance indicators show that our margin has declined because input costs went up sharply this year"
        mock_llm.SHARPEN_REPLIES[:] = [long_one]                    # first answer too long, retry is fine
        h = json.loads(self.post("/api/sharpen", {"headline": "KPIs", "body": []})[2])
        self.assertEqual((h["source"], h["headline"]), ("llm", "Margin fell 1.8 pts, driven by input costs"))
        mock_llm.SHARPEN_REPLIES[:] = [long_one] * 3                 # never fixed: built-in rules, never the long one
        h = json.loads(self.post("/api/sharpen", {"headline": "KPIs", "body": []})[2])
        self.assertEqual(h["source"], "rules")
        self.assertLessEqual(len(h["headline"].split()), 15)
        self.assertEqual(h["headline"], "KPIs")
        self.assertFalse(h["meets_rules"])            # R-03: the UI must not call "KPIs" an action title

    def test_export_pptx_valid(self):
        s = json.loads(self.post("/api/storyline", {"brief": BRIEF, "version": "full"})[2])
        st, h, body = self.post("/api/export/pptx", {"brief": BRIEF, "title": s["title"], "slides": s["slides"]})
        self.assertEqual(st, 200)
        self.assertIn("attachment", h["Content-Disposition"])
        zf = zipfile.ZipFile(io.BytesIO(body))
        names = zf.namelist()
        self.assertIn("ppt/presentation.xml", names)
        self.assertEqual(len([n for n in names if n.startswith("ppt/slides/slide")]), len(s["slides"]))
        import xml.etree.ElementTree as ET
        for n in names:
            if n.endswith(".xml") or n.endswith(".rels"):
                ET.fromstring(zf.read(n))  # every part is well-formed

    def test_export_keeps_overflow_and_no_placeholder_footer(self):  # D-14
        from pptx import Presentation
        bullets = ["Point %d costs SAR %d" % (k, k * 10) for k in range(1, 8)]
        slides = [{"kind": "summary", "headline": "Seven points that matter", "body": bullets, "evidence": {"status": "needed"}}]
        st, h, body = self.post("/api/export/pptx", {"brief": BRIEF, "title": "T", "slides": slides})
        self.assertEqual(st, 200, body)
        self.assertIn("moved to the speaker notes", h["X-Studio-Notice"])
        s = Presentation(io.BytesIO(body)).slides[0]
        everything = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) + s.notes_slide.notes_text_frame.text
        for b in bullets:
            self.assertIn(b, everything)                                  # nothing lost
        self.assertNotIn("Evidence needed", everything)                   # no "Evidence needed: Evidence needed"

    def test_briefing(self):
        s = json.loads(self.post("/api/storyline", {"brief": BRIEF})[2])
        st, _, body = self.post("/api/export/briefing", {"brief": BRIEF, "title": "T", "slides": s["slides"],
                                                         "include": {"brief": True, "notes": True, "qa": True},
                                                         "qa": [{"question": "Q?", "answer": "A", "evidence": "E"}]})
        text = body.decode()
        self.assertIn("# Executive brief", text)
        self.assertIn("# Speaker notes", text)
        self.assertIn("Q?", text)

    def test_extract_docx_and_errors(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("word/document.xml",
                        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                        "<w:p><w:r><w:t>Revenue grew 4%</w:t></w:r></w:p></w:body></w:document>")
        st, _, body = self.post("/api/extract", raw=buf.getvalue(),
                                headers={"Content-Type": "application/octet-stream", "X-Filename": "notes.docx"})
        self.assertEqual(st, 200, body)
        self.assertIn("Revenue grew 4%", json.loads(body)["text"])
        st, _, body = self.post("/api/extract", raw=b"%PDF-1.4 broken", headers={"X-Filename": "x.pdf"})
        self.assertEqual(st, 400)   # unreadable PDF, or PDF support not installed: both are a clear 400
        st, _, body = self.post("/api/extract", raw=b"x", headers={"X-Filename": "old.ppt"})
        self.assertIn("Older Office formats", json.loads(body)["error"])

    def test_bad_requests(self):
        self.assertEqual(self.post("/api/storyline", raw=b"not json")[0], 400)
        self.assertEqual(self.post("/api/storyline", {"brief": dict(BRIEF, topic="")})[0], 400)
        self.assertEqual(self.post("/api/nope", {})[0], 404)


try:
    import pptx  # noqa: F401
    import sample_deck
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False


@unittest.skipUnless(HAVE_PPTX, "python-pptx not installed")
class Decks(Base):
    llm_up = True

    def _upload(self):
        st, _, body = self.post("/api/deck/upload", raw=sample_deck.build(),
                                headers={"Content-Type": "application/octet-stream", "X-Filename": "ops%20review.pptx"})
        self.assertEqual(st, 200, body)
        return json.loads(body)

    def test_upload_inventory(self):
        d = self._upload()
        self.assertEqual(len(d["slides"]), 4)
        self.assertTrue(d["deck_id"])
        s3 = d["slides"][2]
        self.assertEqual((s3["charts"], s3["pictures"]), (1, 1))
        self.assertIn("handling time", d["text"])

    def _open(self, fid):
        from pptx import Presentation
        st, h, body = self.get("/api/download/" + fid)
        self.assertEqual(st, 200)
        self.assertIn("presentationml", h["Content-Type"])
        return Presentation(io.BytesIO(body))

    def test_transform_in_place(self):
        d = self._upload()
        st, _, body = self.post("/api/deck/transform", {"deck_id": d["deck_id"], "selected": [2],
                                "actions": {"headlines": True, "simplify": True}, "brief": BRIEF})
        self.assertEqual(st, 200, body)
        r = json.loads(body)
        self.assertEqual(r["source"], "llm")
        prs = self._open(r["file_id"])
        self.assertEqual(len(prs.slides), 4)                      # all slides kept
        self.assertTrue(prs.slides[1].shapes.title.text.endswith("the key finding"))
        self.assertEqual(prs.slides[0].shapes.title.text, "Quarterly operations review")  # unselected untouched
        body_tf = prs.slides[1].placeholders[1].text_frame
        self.assertEqual(len(body_tf.paragraphs), 3)

    def test_transform_retemplate_carries_visuals(self):
        d = self._upload()
        st, _, body = self.post("/api/deck/transform", {"deck_id": d["deck_id"], "selected": [1, 3, 4],
                                "actions": {"retemplate": True}, "brief": BRIEF})
        self.assertEqual(st, 200, body)
        r = json.loads(body)
        prs = self._open(r["file_id"])
        self.assertEqual(len(prs.slides), 3)
        kinds = [[sh.has_chart if hasattr(sh, "has_chart") else False for sh in s.shapes] for s in prs.slides]
        self.assertTrue(any(kinds[1]))                               # chart rebuilt natively
        self.assertTrue(any(getattr(sh, "has_table", False) and sh.has_table for sh in prs.slides[2].shapes))
        self.assertTrue(any("Picture carried over" in c for c in r["report"][1]["changes"]))

    def test_storyline_chart_export_is_native(self):
        from pptx import Presentation
        slides = [{"kind": "bars", "headline": "Revenue rose 8% on enterprise demand", "body": ["Enterprise +12%"],
                   "chart": {"type": "column", "categories": ["Q1", "Q2"], "series": [{"name": "Rev", "values": ["100", "108"]}]},
                   "evidence": {"status": "linked", "source": "Finance"}}]
        st, _, body = self.post("/api/export/pptx", {"brief": BRIEF, "title": "T", "slides": slides})
        self.assertEqual(st, 200)
        prs = Presentation(io.BytesIO(body))
        self.assertTrue(any(getattr(sh, "has_chart", False) and sh.has_chart for sh in prs.slides[0].shapes))
        self.assertIn("Finance", prs.slides[0].notes_slide.notes_text_frame.text)

    def test_download_messages(self):  # D-12
        for fid, expect in (("nonsense", "File not found"), ("0" * 32, "no longer on the server")):
            with self.assertRaises(urllib.error.HTTPError) as cm:
                self.get("/api/download/" + fid)
            self.assertIn(expect, json.loads(cm.exception.read())["error"])
        st, _, body = self.post("/api/extract", raw=b"MZ\x90\x00", headers={"X-Filename": "tool.exe"})
        self.assertIn(".pdf", json.loads(body)["error"])                  # PDF is supported here, so it is listed

    def test_expired_deck(self):
        st, _, body = self.post("/api/deck/transform", {"deck_id": "0" * 32, "selected": [1], "actions": {"headlines": True}, "brief": BRIEF})
        self.assertEqual(st, 400)
        self.assertIn("Upload it again", json.loads(body)["error"])

    # ---- N-01 / N-08: re-templating a deck this app exported (badges, tiles, footers)
    def _upload_app_deck(self):
        slides = [{"kind": "title", "headline": "Logistics update", "body": ["Operations"]},
                  {"kind": "summary", "headline": "Overview", "body": ["On-time delivery 91% vs 95% target", "Warehouse B reached 97%",
                   "Carrier X contract runs to March 2027", "Fuel surcharge up 6%", "Returns backlog 1,800 parcels"]},
                  {"kind": "kpi", "headline": "KPIs", "body": ["On-time 91% vs 95%", "Cost per parcel EUR 4.10 vs 3.90", "Damage rate 0.8% vs 0.5%"]},
                  {"kind": "roadmap", "headline": "Next steps", "body": ["Discuss", "Review", "Follow up"]}]
        st, _, exported = self.post("/api/export/pptx", {"brief": BRIEF, "title": "Logistics update", "slides": slides})
        self.assertEqual(st, 200)
        st, _, body = self.post("/api/deck/upload", raw=exported, headers={"X-Filename": "Logistics_update_test.pptx"})
        return json.loads(body)

    def test_retemplate_app_deck_no_duplicates(self):
        d = self._upload_app_deck()
        for actions in ({"headlines": True, "simplify": True, "retemplate": True}, {"retemplate": True}):
            st, _, body = self.post("/api/deck/transform", {"deck_id": d["deck_id"], "selected": [2, 3, 4], "actions": actions, "brief": BRIEF})
            self.assertEqual(st, 200, body)
            r = json.loads(body)
            prs = self._open(r["file_id"])
            for s in prs.slides:
                title_id = s.shapes.title.shape_id if s.shapes.title is not None else None
                paras = [p.text.strip() for sh in s.shapes if sh.has_text_frame and sh.shape_id != title_id
                         for p in sh.text_frame.paragraphs if p.text.strip()]
                content = [p for p in paras if not re.fullmatch(r"\d+ / \d+|Corporate Master", p)]
                self.assertFalse([p for p in content if re.fullmatch(r"\d{1,2}", p)], (actions, content))   # no stray badges
                self.assertEqual(len(content), len(set(content)), (actions, content))                      # nothing twice
                if actions.get("simplify"):
                    self.assertLessEqual(len(content), 5, content)
                notes = s.notes_slide.notes_text_frame.text
                self.assertNotIn("Small print", notes)                  # old footer and page numbers are not content
            joined = lambda s: " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)  # noqa: E731
            for fact in ("91%", "97%", "March 2027", "6%", "1,800"):
                self.assertIn(fact, joined(prs.slides[0]), actions)
            for fact in ("4.10", "3.90", "0.8%", "0.5%"):
                self.assertIn(fact, joined(prs.slides[1]), actions)
            steps = joined(prs.slides[2])
            self.assertLess(steps.index("Discuss"), steps.index("Review"), steps)  # order kept (N-08)
            self.assertLess(steps.index("Review"), steps.index("Follow up"), steps)

    def test_customer_deck_groups_smartart_tables(self):  # round 3: a deck built in PowerPoint, not by this app
        st, _, body = self.post("/api/deck/upload", raw=sample_deck.build_customer_deck(), headers={"X-Filename": "customer.pptx"})
        self.assertEqual(st, 200, body)
        d = json.loads(body)
        facts = ["93.5%", "91.2%", "7.8", "8.5", "EUR 2.4m", "OTIF 96%", "82%", "Asia", "Leeds", "97%", "S&OP", "12 SKUs",
                 "3PL", "Chief Operating Officer", "December", "January", "Q1 2027", "Executive Committee", "October 2026"]
        for actions in ({"retemplate": True}, {"headlines": True, "simplify": True, "retemplate": True}, {"headlines": True, "simplify": True}):
            st, _, body = self.post("/api/deck/transform", {"deck_id": d["deck_id"], "selected": [1, 2, 3, 4, 5, 6],
                                    "actions": actions, "brief": BRIEF})
            self.assertEqual(st, 200, body)
            r = json.loads(body)
            prs = self._open(r["file_id"])
            from studio.slides.deck import box_text                       # reads inside groups and SmartArt too
            text = " ".join(p for s in prs.slides for sh in s.shapes for p in box_text(sh))
            text += " ".join(s.notes_slide.notes_text_frame.text for s in prs.slides)
            self.assertEqual([f for f in facts if f not in text], [], actions)
            for s in prs.slides:                                             # nothing on a slide twice
                paras = [p.text.strip() for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs if p.text.strip()]
                paras = [p for p in paras if not re.fullmatch(r"\d+ / \d+|Corporate Master", p)]
                self.assertEqual(len(paras), len(set(paras)), (actions, paras))
            last = prs.slides[5]
            self.assertTrue(any(getattr(sh, "has_table", False) and sh.has_table for sh in last.shapes))
            self.assertTrue(any(getattr(sh, "has_chart", False) and sh.has_chart for sh in last.shapes))
            if actions.get("retemplate"):
                warn = " ".join(w for x in r["report"] for w in x["warnings"])
                self.assertIn("group's text was carried over", warn)
                self.assertIn("SmartArt diagram's text was carried over", warn)

    def test_deck_saved_by_powerpoint(self):  # round 4: a deck built and saved by PowerPoint itself
        raw = (ROOT / "tests" / "fixtures" / "powerpoint_northwind.pptx").read_bytes()
        from studio.slides import deck as deck_ops
        texts = {t["index"]: t for t in deck_ops.slide_texts(raw, set(range(1, 7)), whole_slide=True)}
        self.assertIn("34.5", texts[3]["data"][0])                              # S-02: chart values reach the model
        self.assertIn("EUR 71.2m", texts[4]["data"][0])                         # ... and table cells
        self.assertTrue(texts[1]["cover"])
        # S-03: a headline written from the slide's sentence does not reappear as a point under it
        from studio import catalog
        from studio import slides as slide_engine
        tpl = catalog.template("corporate")
        out, _rep = deck_ops.retemplate(raw, {5}, {5: {"title": "Lead time from order to shelf averages 19 days", "body": None}},
                                        tpl, {"template_name": tpl["name"], "date": "x", "title": "x"}, slide_engine.template_path(tpl))
        lines = [p for sh in deck_ops.open_deck(out).slides[0].shapes for p in deck_ops.box_text(sh)]
        self.assertEqual(sum("averages 19 days" in p for p in lines), 1, lines)
        self.assertIn("On shelf after 12 days", lines)
        st, _, body = self.post("/api/deck/upload", raw=raw, headers={"X-Filename": "Northwind.pptx"})
        d = json.loads(body)
        facts = ["184.6m", "6.2%", "31%", "38.9%", "1.4 pts", "0.8 pts", "0.6 pts", "74", "65", "2 days", "5 days",
                 "12 days", "19 days", "3 stores", "11%", "Utrecht", "Executive Committee", "October 2026"]
        for actions in ({"retemplate": True}, {"headlines": True, "simplify": True, "retemplate": True}, {"headlines": True, "simplify": True}):
            st, _, body = self.post("/api/deck/transform", {"deck_id": d["deck_id"], "selected": [1, 2, 3, 4, 5, 6],
                                    "actions": actions, "brief": BRIEF})
            self.assertEqual(st, 200, body)
            r = json.loads(body)
            prs = self._open(r["file_id"])
            text = " ".join(p for s in prs.slides for sh in s.shapes for p in deck_ops.box_text(sh))
            notes = " ".join(s.notes_slide.notes_text_frame.text for s in prs.slides)
            if not actions.get("retemplate"):   # in place, SmartArt text stays in its own diagram part
                text += " ".join(deck_ops.slide_texts(raw, {5}, whole_slide=True)[0]["body"])
            self.assertEqual([f for f in facts if f not in text + notes], [], actions)
            for phrase in ("Lead with sales growth", "Online keeps gaining share", "roll-out in Q1 2027"):
                self.assertIn(phrase, notes, actions)                           # speaker notes survive
            shapes = lambda i: prs.slides[i].shapes  # noqa: E731
            self.assertTrue(any(getattr(sh, "has_chart", False) and sh.has_chart for sh in shapes(2)), actions)
            self.assertTrue(any(getattr(sh, "has_table", False) and sh.has_table for sh in shapes(3)), actions)
            self.assertTrue(any(sh.shape_type == 13 for sh in shapes(5)), actions)  # the picture
            if actions.get("retemplate"):
                for s in prs.slides:                                            # S-03: headline never repeated below
                    head = s.shapes.title.text_frame.text.strip().rstrip(".").lower() if s.shapes.title is not None else ""
                    others = [p.rstrip(".").lower() for sh in s.shapes if sh.shape_id != getattr(s.shapes.title, "shape_id", None)
                              for p in deck_ops.box_text(sh)]
                    self.assertNotIn(head, others, actions)

    def test_in_place_never_copies_other_boxes(self):  # N-01 in place: a KPI tile as the "main text"
        from studio.checks import rewrite_problem
        original = {"title": "Key metrics", "body": ["94%", "On-time delivery"],
                    "other_text": ["3.2 days – Average lead time", "SAR 1.4m – Monthly freight spend"]}
        copy = {"title": "", "body": ["94% on-time delivery", "3.2 days – Average lead time", "SAR 1.4m – Monthly freight spend"]}
        self.assertIn("other boxes", rewrite_problem(original, copy))
        d = self._upload_report_deck()
        st, _, body = self.post("/api/deck/transform", {"deck_id": d["deck_id"], "selected": [3],
                                "actions": {"simplify": True}, "brief": BRIEF})
        r = json.loads(body)
        s = self._open(r["file_id"]).slides[2]
        paras = [p.text for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs if p.text.strip()]
        self.assertEqual(len(paras), len(set(paras)), paras)                  # each tile's text once
        self.assertNotIn("Main text simplified", " ".join(r["report"][0]["changes"]))

    def test_transform_title_direction_checked(self):  # N-03
        from studio.checks import rewrite_problem
        original = {"title": "KPIs", "body": ["On-time 91% vs 95%", "Cost per parcel EUR 4.10 vs 3.90", "Damage rate 0.8% vs 0.5%"]}
        bad = {"title": "Cost per parcel, on-time delivery, and damage rate are all below target", "body": None}
        self.assertIn("above target", rewrite_problem(original, bad))
        good = {"title": "All three KPIs miss target: cost 4.10 vs 3.90, damage 0.8% vs 0.5%", "body": None}
        self.assertEqual(rewrite_problem(original, good), "")

    # ---- D-02: board assessment judges the slides, not the brief
    def test_board_pres08_fails_without_decision_slide(self):
        slides = [{"kind": "summary", "headline": "Margin fell 1.8 pts on input costs", "body": ["Freight +0.9 pts"]},
                  {"kind": "roadmap", "headline": "Repricing lands in Q1", "body": ["Pilot in November"]}]
        st, _, body = self.post("/api/board-assessment", {"brief": BRIEF, "slides": slides, "context": {}})
        self.assertEqual(st, 200, body)
        r = json.loads(body)
        self.assertEqual(r["source"], "llm")
        pres08 = next(c for c in r["compliance"] if c["id"] == "PRES-08")
        self.assertEqual(pres08["result"], "FAIL")
        clarity = next(d for d in r["board_readiness"] if d["id"] == "decision_clarity")
        self.assertLessEqual(clarity["score"], 2)
        self.assertIn("appears only in the brief", pres08["evidence"])
        # R-02: a failing critical check caps the overall readiness score
        self.assertEqual(r["readiness"]["critical_fails"], ["PRES-08"])
        self.assertFalse(r["readiness"]["board_ready"])
        self.assertTrue(r["readiness"]["capped"])
        self.assertEqual(r["readiness"]["score"], 2.5)
        self.assertGreater(r["readiness"]["raw_score"], 2.5)
        # S-01: the summary leads with the verdict and keeps none of the model's "Ready."
        self.assertTrue(r["summary"].startswith("Not board-ready: critical check PRES-08 fails"), r["summary"])
        self.assertNotRegex(r["summary"].split(").", 1)[1], r"(?i)\bready\b")
        self.assertEqual(r["summary"], (r["verdict"] + " " + r["summary_detail"]).strip())   # the page shows the verdict once
        self.assertNotIn("Not board-ready", r["summary_detail"])
        # N-04: the ask is on slide 1 but the deck closes on a roadmap: still FAIL, and the wording says so
        opener = dict(slides[0], body=["Approve a 3% mid-market price increase from January", "Freight +0.9 pts"])
        r = json.loads(self.post("/api/board-assessment", {"brief": BRIEF, "slides": [opener, slides[1]], "context": {}})[2])
        pres08 = next(c for c in r["compliance"] if c["id"] == "PRES-08")
        self.assertEqual(pres08["result"], "FAIL")
        self.assertIn("stated on slide 1", pres08["evidence"])
        self.assertEqual(next(d for d in r["board_readiness"] if d["id"] == "decision_clarity")["score"], 3)
        slides.append({"kind": "decision", "headline": "Approve a 3% price increase", "body": ["Approve today"]})
        r = json.loads(self.post("/api/board-assessment", {"brief": BRIEF, "slides": slides, "context": {}})[2])
        self.assertEqual(next(c for c in r["compliance"] if c["id"] == "PRES-08")["result"], "PASS")
        self.assertFalse(r["readiness"]["capped"])                       # no critical failure: no cap
        self.assertGreater(r["readiness"]["score"], 2.5)
        # round 7: a passing deck never reads "not board-ready because it fails the critical PRES-07"
        self.assertTrue(r["summary"].startswith("Board-ready: no critical check fails."), r["summary"])
        self.assertNotIn("not board-ready", r["summary"].lower())
        self.assertNotIn("critical PRES-07", r["summary"])
        self.assertIn("Message titles are crisp.", r["summary"])               # the model's other sentences stay

    def test_board_pres03_decided_from_slides(self):  # round 6/7: PRES-03 failed by mistake ~1 run in 10
        from studio.services import _enforce_opening_ask
        def run(slides, model_says, decision="Approve a 3% price increase"):
            result = {"compliance": [{"id": "PRES-03", "result": model_says, "evidence": "", "action": ""}]}
            _enforce_opening_ask(result, slides, decision)
            return result["compliance"][0]["result"]
        opener = [{"kind": "title", "headline": "Q3 review", "body": []},
                  {"kind": "summary", "headline": "Approve a 3% price increase to defend margin", "body": ["Margin down 2.1 pts"]}]
        self.assertEqual(run(opener, "FAIL"), "PASS")                          # the model's false FAIL is corrected
        late = [{"kind": "title", "headline": "Q3 review", "body": []}, {"kind": "kpi", "headline": "Margin fell 2.1 pts", "body": ["Freight"]},
                {"kind": "decision", "headline": "Approve a 3% price increase", "body": []}]
        self.assertEqual(run(late, "PASS"), "FAIL")                            # ask only at the end: not an opening ask
        self.assertEqual(run(late, "PASS", decision=""), "PASS")               # no decision in the brief: model judges
        info = [{"kind": "summary", "headline": "For information: safety incidents fell 12%", "body": []}]
        self.assertEqual(run(info, "FAIL", decision=""), "PASS")

    def test_board_rejects_meeting_before_submission(self):  # D-07
        slides = [{"kind": "decision", "headline": "Approve a 3% price increase", "body": ["Approve today"]}]
        st, _, body = self.post("/api/board-assessment", {"brief": BRIEF, "slides": slides,
                                "context": {"submission_date": "2026-11-10", "meeting_date": "2026-11-02"}})
        self.assertEqual(st, 400)
        self.assertIn("before the submission date", json.loads(body)["error"])
        st, _, body = self.post("/api/board-assessment", {"brief": BRIEF, "slides": slides,
                                "context": {"submission_date": "2026-11-10", "meeting_date": "next week"}})
        self.assertEqual(st, 400)
        r = json.loads(self.post("/api/board-assessment", {"brief": BRIEF, "slides": slides,
                                 "context": {"submission_date": "2026-11-10", "meeting_date": "2026-11-10"}})[2])
        self.assertEqual(r["compliance"][0]["result"], "FAIL")                 # same day: 0 business days

    # ---- D-01: transform must never lose content
    def _upload_report_deck(self):
        st, _, body = self.post("/api/deck/upload", raw=sample_deck.build_report_deck(), headers={"X-Filename": "logistics.pptx"})
        self.assertEqual(st, 200, body)
        return json.loads(body)

    def test_transform_rejects_lossy_rewrite(self):
        d = self._upload_report_deck()
        mock_llm.TRANSFORM_MODE, mock_llm.TRANSFORM_CALLS[:] = "drop", []
        try:
            st, _, body = self.post("/api/deck/transform", {"deck_id": d["deck_id"], "selected": [2],
                                    "actions": {"simplify": True}, "brief": BRIEF})
        finally:
            mock_llm.TRANSFORM_MODE = "merge"
        self.assertEqual(st, 200, body)
        r = json.loads(body)
        self.assertEqual(len(mock_llm.TRANSFORM_CALLS), 2)              # asked once, re-asked once
        self.assertTrue(any("Left unchanged" in w and "SAR" not in w for w in r["report"][0]["warnings"]), r["report"])
        self.assertNotIn("Main text simplified", " ".join(r["report"][0]["changes"]))
        text = [p.text for p in self._open(r["file_id"]).slides[1].placeholders[1].text_frame.paragraphs]
        self.assertEqual(len(text), 5)                                  # all five bullets still there
        self.assertIn("Contract with Bahri renews in Q1", text)

    def test_transform_keeps_text_boxes_and_ignores_footer(self):
        d = self._upload_report_deck()
        self.assertEqual(d["slides"][3]["body"], ["Tender by 30 Nov"])   # the footer is not the body
        st, _, body = self.post("/api/deck/transform", {"deck_id": d["deck_id"], "selected": [2, 3, 4],
                                "actions": {"headlines": True, "simplify": True, "retemplate": True}, "brief": BRIEF})
        self.assertEqual(st, 200, body)
        r = json.loads(body)
        prs = self._open(r["file_id"])
        text = lambda s: " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)  # noqa: E731
        for fact in ("94%", "3.2 days", "SAR 1.4m", "Average lead time"):
            self.assertIn(fact, text(prs.slides[1]))                   # every KPI tile carried over
        self.assertIn("SAR 410", text(prs.slides[0]))
        self.assertIn("Contract with Bahri", text(prs.slides[0]))
        self.assertIn("Tender by 30 Nov", text(prs.slides[2]))
        self.assertNotIn("Evidence needed: Evidence needed", text(prs.slides[2]))
        self.assertIn("Evidence needed", prs.slides[2].notes_slide.notes_text_frame.text)  # small print kept in notes


class WithoutLLM(Base):
    llm_up = False

    def test_fallback_deck_rewrite(self):
        if not HAVE_PPTX:
            self.skipTest("python-pptx not installed")
        st, _, body = self.post("/api/deck/upload", raw=sample_deck.build(), headers={"X-Filename": "d.pptx"})
        d = json.loads(body)
        r = json.loads(self.post("/api/deck/transform", {"deck_id": d["deck_id"], "selected": [2],
                                 "actions": {"headlines": True, "simplify": True}, "brief": BRIEF})[2])
        self.assertEqual(r["source"], "rules")
        self.assertEqual(r["report"][0]["after"], "Customer service")  # "Overview of" stripped by the rules

    def test_fallback_storyline(self):
        st, _, body = self.post("/api/storyline", {"brief": BRIEF, "version": "summary"})
        s = json.loads(body)
        self.assertEqual(s["source"], "rules")
        self.assertIn("warning", s)
        self.assertEqual(len(s["slides"]), 5)
        a = json.loads(self.post("/api/advise", {"message": "compare three vendors"})[2])
        self.assertEqual(a["primary"]["kind"], "matrix")


if __name__ == "__main__":
    unittest.main()
