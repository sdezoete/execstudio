"""Builds a small "existing" deck with a picture, a table and a chart, for tests and demos.

    python tests/sample_deck.py sample.pptx
"""
import io
import struct
import sys
import zlib

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def _png(w=320, h=180, rgb=(31, 71, 201)):
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))

    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def build():
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Quarterly operations review"
    s1.placeholders[1].text = "Operations team"

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Overview of customer service"
    body = s2.placeholders[1].text_frame
    body.text = "Average handling time went down from 7.2 to 6.1 minutes over the quarter"
    for t in ["First contact resolution improved to 78%", "Staffing was below plan in August",
              "Customer satisfaction stable at 4.3", "New ticketing tool live in two sites",
              "Training completed for 85% of agents", "Backlog cleared by mid-September"]:
        body.add_paragraph().text = t
    s2.notes_slide.notes_text_frame.text = "Mention the August staffing gap."

    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Ticket volumes"
    cd = CategoryChartData()
    cd.categories = ["Jul", "Aug", "Sep"]
    cd.add_series("Tickets", (1200, 1350, 1100))
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.8), Inches(5), Inches(4), cd)
    s3.shapes.add_picture(io.BytesIO(_png()), Inches(6.3), Inches(2), Inches(3.2))

    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "Site comparison"
    tbl = s4.shapes.add_table(3, 3, Inches(1), Inches(2), Inches(8), Inches(1.5)).table
    for r, row in enumerate([["Site", "AHT", "FCR"], ["Riyadh", "5.9", "80%"], ["Dammam", "6.4", "75%"]]):
        for c, v in enumerate(row):
            tbl.cell(r, c).text = v
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_report_deck():
    """The shapes behind the 23 Oct content-loss defect (D-01): five bullets with figures, three KPI
    tiles in separate text boxes, and a short "Next steps" slide whose longest text is a footer."""
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Logistics update"
    s1.placeholders[1].text = "Supply chain team"

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Overview of freight costs"
    body = s2.placeholders[1].text_frame
    body.text = "Freight cost per tonne rose 12% to SAR 410"
    for t in ["Diesel surcharge added 0.9 pts", "Port fees in Jeddah up 6%", "Two carriers exited in July",
              "Contract with Bahri renews in Q1"]:
        body.add_paragraph().text = t

    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Key metrics"
    for k, (num, label) in enumerate([("94%", "On-time delivery"), ("3.2 days", "Average lead time"),
                                      ("SAR 1.4m", "Monthly freight spend")]):
        tf = s3.shapes.add_textbox(Inches(0.6 + k * 3.1), Inches(2.2), Inches(2.8), Inches(1.6)).text_frame
        tf.text = num
        tf.add_paragraph().text = label

    s4 = prs.slides.add_slide(prs.slide_layouts[1])
    s4.shapes.title.text = "Next steps"
    s4.placeholders[1].text = "Tender by 30 Nov"
    s4.shapes.add_textbox(Inches(0.5), Inches(6.65), Inches(9), Inches(0.35)).text_frame.text = "Evidence needed: Evidence needed"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


_DGM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _add_smartart(slide, nodes, x, y, w, h):
    """A SmartArt graphic frame plus its diagram-data part, which is where PowerPoint keeps the text.
    Only the data part is built (no layout/style parts), which is all a text reader needs."""
    from lxml import etree
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    pts = ['<dgm:pt modelId="0" type="doc"><dgm:t><a:bodyPr/><a:p><a:endParaRPr/></a:p></dgm:t></dgm:pt>']
    for k, text in enumerate(nodes, 1):
        pts.append('<dgm:pt modelId="%d"><dgm:t><a:bodyPr/><a:p><a:r><a:t>%s</a:t></a:r></a:p></dgm:t></dgm:pt>' % (k, text))
        pts.append('<dgm:pt modelId="%d" type="parTrans"><dgm:t><a:bodyPr/><a:p><a:r><a:t>connector</a:t></a:r></a:p></dgm:t></dgm:pt>' % (100 + k))
    xml = ('<dgm:dataModel xmlns:dgm="%s" xmlns:a="%s"><dgm:ptLst>%s</dgm:ptLst><dgm:cxnLst/></dgm:dataModel>'
           % (_DGM_NS, _A_NS, "".join(pts)))
    part = Part(PackURI("/ppt/diagrams/data%d.xml" % slide.slide_id), "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
                slide.part.package, xml.encode("utf-8"))
    rid = slide.part.relate_to(part, _R_NS + "/diagramData")
    frame = etree.fromstring(
        '<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="%s" xmlns:r="%s">'
        '<p:nvGraphicFramePr><p:cNvPr id="%d" name="Diagram"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
        '<p:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></p:xfrm><a:graphic><a:graphicData uri="%s">'
        '<dgm:relIds xmlns:dgm="%s" r:dm="%s" r:lo="" r:qs="" r:cs=""/></a:graphicData></a:graphic></p:graphicFrame>'
        % (_A_NS, _R_NS, 900 + slide.slide_id, Inches(x), Inches(y), Inches(w), Inches(h), _DGM_NS, _DGM_NS, rid))
    slide.shapes._spTree.append(frame)  # noqa: SLF001


def build_customer_deck():
    """A deck built the way people build them in PowerPoint, not exported by this app: multi-level
    bullets, a two-content layout, a grouped process diagram with a caption outside the group, a
    SmartArt diagram, a table with a native chart, and speaker notes."""
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "FY26 supply chain review"
    s1.placeholders[1].text = "Prepared for the Executive Committee"
    s1.placeholders[1].text_frame.add_paragraph().text = "October 2026"

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Where we stand"
    tf = s2.placeholders[1].text_frame
    tf.text = "Service level reached 93.5% in Q3"
    for text, level in [("Up from 91.2% in Q2", 1), ("Inventory turns 7.8 vs 8.5 target", 0),
                        ("Northern DC below 6 turns", 1), ("Expedited freight EUR 2.4m year to date", 0)]:
        p = tf.add_paragraph()
        p.text, p.level = text, level
    s2.notes_slide.notes_text_frame.text = "Stress the service gain before the inventory gap."

    s3 = prs.slides.add_slide(prs.slide_layouts[3])  # Two Content
    s3.shapes.title.text = "Strengths and gaps"
    s3.placeholders[1].text = "Supplier OTIF 96%"
    s3.placeholders[1].text_frame.add_paragraph().text = "Forecast accuracy 82%"
    s3.placeholders[2].text = "Two single-source SKUs in Asia"
    s3.placeholders[2].text_frame.add_paragraph().text = "Warehouse Leeds at 97% capacity"

    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "Operating model"
    grp = s4.shapes.add_group_shape()
    for k, text in enumerate(["Plan: monthly S&OP with Finance", "Source: dual-source 12 SKUs", "Deliver: two new 3PL lanes"]):
        grp.shapes.add_textbox(Inches(0.5 + k * 3.1), Inches(2.4), Inches(2.9), Inches(1.2)).text_frame.text = text
    s4.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(0.6)).text_frame.text = "Owner: Chief Operating Officer"

    s5 = prs.slides.add_slide(prs.slide_layouts[5])
    s5.shapes.title.text = "Roadmap"
    _add_smartart(s5, ["Assess network by December", "Select partners in January", "Launch pilot in Q1 2027"], 0.5, 2.0, 9, 3)

    s6 = prs.slides.add_slide(prs.slide_layouts[5])
    s6.shapes.title.text = "Cost by region"
    tbl = s6.shapes.add_table(3, 3, Inches(0.5), Inches(1.8), Inches(4.5), Inches(1.4)).table
    for r, row in enumerate([["Region", "Cost", "Change"], ["North", "EUR 5.1m", "+4%"], ["South", "EUR 3.7m", "-2%"]]):
        for c, v in enumerate(row):
            tbl.cell(r, c).text = v
    cd = CategoryChartData()
    cd.categories = ["North", "South"]
    cd.add_series("Cost (EUR m)", (5.1, 3.7))
    s6.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5.3), Inches(1.8), Inches(4.2), Inches(3.5), cd)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    with open(sys.argv[1] if len(sys.argv) > 1 else "sample.pptx", "wb") as fh:
        fh.write(build())
