"""Working with existing decks (requires python-pptx).

  inventory(data)                       list what is on every slide
  slide_texts(data, indices)            title + body text per slide, for the LLM
  apply_in_place(data, changes)         rewrite titles/body text, keep everything else
  retemplate(data, indices, changes, tpl, meta, template_path)
                                        rebuild selected slides on the corporate template,
                                        carrying pictures, tables and charts across
"""
import copy
import io
import re

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from . import layouts
from . import pptx_engine as eng
from .pptx_engine import inches

TITLE_TYPES = eng.TITLE_TYPES
CHROME_TYPES = (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE)


class DeckError(Exception):
    pass


def open_deck(data):
    try:
        return Presentation(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001  (python-pptx raises several types for bad files)
        raise DeckError("This file could not be opened as a PowerPoint deck. Save it as .pptx and try again.") from e


# ------------------------------------------------------------------ shape helpers
def _walk(shapes):
    for shp in shapes:
        yield shp
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shp.shapes)


def _text_of(shp):
    if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
        return [p.text.strip() for p in shp.text_frame.paragraphs if p.text.strip()]
    return []


def _has_text(shp):
    return getattr(shp, "has_text_frame", False) and shp.has_text_frame and bool(shp.text_frame.text.strip())


_DGM = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"


def _smartart_text(shp):
    """The text of a SmartArt diagram. It lives in the diagram's data part, not on the slide, so
    python-pptx shows the shape as an empty graphic frame."""
    rel = shp._element.find(".//" + _DGM + "relIds")  # noqa: SLF001
    if rel is None:
        return []
    try:
        root = etree.fromstring(shp.part.related_part(rel.get(qn("r:dm"))).blob)
    except Exception:  # noqa: BLE001  (a damaged or unusual diagram: carry what we can, never fail the deck)
        return []
    out = []
    for pt in root.iter(_DGM + "pt"):
        if pt.get("type", "node") == "node":  # skip the document node and connector "transitions"
            t = "".join(x.text or "" for x in pt.iter(qn("a:t"))).strip()
            if t:
                out.append(t)
    return out


def box_text(shp):
    """The paragraphs a shape shows: a text box's own, every shape's inside a group (in reading order),
    or a SmartArt diagram's. Groups and SmartArt are how real decks hold process steps and callouts."""
    if _has_text(shp):
        return _text_of(shp)
    if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [p for child in reading_order(list(shp.shapes)) for p in box_text(child)]
    if shp.element.tag == qn("p:graphicFrame"):
        return _smartart_text(shp)
    return []


def _slide_height(slide):
    try:
        return slide.part.package.presentation_part.presentation.slide_height
    except Exception:  # noqa: BLE001
        return None


def is_chrome(shp, slide_h):
    """Footer, date and slide-number placeholders, and short text along the bottom edge (a source line,
    a confidentiality mark, or this studio's own evidence note). Never a slide's title or main text."""
    if shp.is_placeholder and shp.placeholder_format.type in CHROME_TYPES:
        return True
    return bool(slide_h) and shp.top is not None and shp.top >= 0.85 * slide_h and len(" ".join(box_text(shp))) <= 160


def find_title(slide):
    for shp in slide.shapes:
        if shp.is_placeholder and shp.placeholder_format.type in TITLE_TYPES and shp.has_text_frame:
            return shp
    # no title placeholder: the highest short text box in the top quarter of the slide
    slide_h = _slide_height(slide)
    cands = [s for s in slide.shapes if _has_text(s) and len(s.text_frame.text.strip()) <= 160
             and s.top is not None and not is_chrome(s, slide_h)]
    cands.sort(key=lambda s: s.top)
    return cands[0] if cands else None


def find_body(slide, title):
    """The main text container: the non-title shape with the most text, preferring body placeholders.
    Footers and bottom-edge small print are never chosen, however long they are."""
    best, best_score = None, 0
    title_id = title.shape_id if title is not None else None
    slide_h = _slide_height(slide)
    for shp in slide.shapes:
        if shp.shape_id == title_id or not _has_text(shp) or is_chrome(shp, slide_h):
            continue
        n = len(shp.text_frame.text.strip())
        if n < 3:
            continue
        score = n * (1.5 if shp.is_placeholder else 1.0)
        if score > best_score:
            best, best_score = shp, score
    return best


_ROW_TOLERANCE = 304800  # 1/3 inch in EMU: boxes whose tops are this close sit on one row
_PAGE_NUMBER = re.compile(r"^\s*\d{1,3}\s*(/|of)\s*\d{1,3}\s*$", re.I)


def _squash(text):
    """Letters and digits only, lower case: "Lead time averages 19 days." == "lead time averages 19 days"."""
    return re.sub(r"[\W_]+", "", (text or "").lower())


def _is_badge(text):
    """Number badges, bullets and step markers ("1", "02", "•", "A") are decoration, not content."""
    t = text.strip()
    return len(t) <= 2 or bool(re.fullmatch(r"\W*\d{1,2}\W*", t))


def _rows(shapes):
    """Rows top to bottom, each left to right. Sorting on the exact top position scrambles items laid
    out side by side whose tops differ by a hair (test report N-08)."""
    rows = []
    for shp in sorted(shapes, key=lambda s: s.top or 0):
        if rows and abs((shp.top or 0) - (rows[-1][0].top or 0)) <= _ROW_TOLERANCE:
            rows[-1].append(shp)
        else:
            rows.append([shp])
    return [sorted(row, key=lambda s: s.left or 0) for row in rows]


def reading_order(shapes):
    return [shp for row in _rows(shapes) for shp in row]


def _is_tile(shp, paras):
    """A KPI tile: a free text box holding a short figure over its label ("94%" / "On-time delivery").
    Its two lines read as one point; any other box's paragraphs are separate points (a placeholder's
    two bullets stay two bullets)."""
    return (len(paras) == 2 and _has_text(shp) and not shp.is_placeholder
            and len(paras[0]) <= 12 and any(ch.isdigit() for ch in paras[0]))


def _column_heads(row, below):
    """True when `row` is a row of short labels sitting over the same columns as `below`
    ("Phase 1 | Phase 2 | Phase 3" above three items): those read as label – item pairs."""
    return (len(row) == len(below) >= 2
            and all(len(" ".join(box_text(s))) <= 25 for s in row)
            and all(abs((a.left or 0) - (b.left or 0)) <= _ROW_TOLERANCE for a, b in zip(row, below)))


def other_text(slide, title, body):
    """Text in every other box, in reading order: (content, small_print). content is one line per box
    (KPI tiles, callouts, labels), without number badges; small_print is the bottom-edge text that
    is_chrome() skips, minus slide numbers and dates."""
    skip = {s.shape_id for s in (title, body) if s is not None}
    slide_h = _slide_height(slide)
    content, small = [], []
    for shp in reading_order([s for s in slide.shapes if s.shape_id not in skip and box_text(s)]):
        line = " – ".join(box_text(shp))
        if not is_chrome(shp, slide_h):
            if not _is_badge(line):
                content.append(line)
        elif not (shp.is_placeholder and shp.placeholder_format.type in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE)) \
                and not _PAGE_NUMBER.match(line):
            small.append(line)
    return content, small


def slide_content(slide, title):
    """All of a slide's text except its title, as one list in reading order: the main text box's
    paragraphs and every other box, without badges, footers or repeats. This is what a rebuilt slide
    must carry, and what a rewrite of it must preserve."""
    slide_h = _slide_height(slide)
    title_id = title.shape_id if title is not None else None
    boxes = [s for s in slide.shapes if s.shape_id != title_id and box_text(s)
             and not is_chrome(s, slide_h) and not _is_badge(" ".join(box_text(s)))]
    rows = _rows(boxes)
    cover = _is_cover(slide)  # a cover's subtitle and date stamp are separate lines, never one (R-04)
    out, k = [], 0

    def add(line):
        if line and not _is_badge(line) and line not in out:
            out.append(line)

    while k < len(rows):
        if k + 1 < len(rows) and _column_heads(rows[k], rows[k + 1]):
            for head, item in zip(rows[k], rows[k + 1]):
                add(" – ".join(box_text(head)) + " – " + " – ".join(box_text(item)))
            k += 2
            continue
        for shp in rows[k]:
            paras = box_text(shp)
            for line in ([" – ".join(paras)] if not cover and _is_tile(shp, paras) else paras):
                add(line)
        k += 1
    return out


def _visuals(slide):
    pics, charts, tables, other = [], [], [], []
    for shp in slide.shapes:
        st = shp.shape_type
        if getattr(shp, "has_chart", False) and shp.has_chart:
            charts.append(shp)
        elif getattr(shp, "has_table", False) and shp.has_table:
            tables.append(shp)
        elif st == MSO_SHAPE_TYPE.PICTURE or (shp.is_placeholder and hasattr(shp, "image") and _has_image(shp)):
            pics.append(shp)
        elif st in (MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.MEDIA, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
                    MSO_SHAPE_TYPE.LINKED_OLE_OBJECT) or (st == MSO_SHAPE_TYPE.PLACEHOLDER and not shp.has_text_frame):
            other.append(shp)
        elif shp.element.tag == qn("p:graphicFrame"):
            other.append(shp)  # SmartArt and other diagrams
    return pics, charts, tables, other


def _has_image(shp):
    try:
        return bool(shp.image.blob)
    except Exception:  # noqa: BLE001
        return False


# ------------------------------------------------------------------ inventory
def inventory(data):
    prs = open_deck(data)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title = find_title(slide)
        body = find_body(slide, title)
        pics, charts, tables, other = _visuals(slide)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
        out.append({
            "index": i,
            "title": title.text_frame.text.strip() if title else "",
            "body": _text_of(body)[:12] if body else [],
            "layout": slide.slide_layout.name,
            "pictures": len(pics), "charts": len(charts), "tables": len(tables), "other": len(other),
            "has_notes": bool(notes),
        })
    if not out:
        raise DeckError("This deck has no slides.")
    return {"slides": out, "width_in": round(inches(prs.slide_width), 2), "height_in": round(inches(prs.slide_height), 2)}


# ------------------------------------------------------------------ text replacement
def replace_paragraphs(tf, lines):
    """Replace the text in a frame, re-using the first paragraph's and run's formatting for every line."""
    tx = tf._txBody  # noqa: SLF001
    paras = tx.findall(qn("a:p"))
    if not paras:
        tf.text = "\n".join(lines)
        return
    proto = copy.deepcopy(paras[0])
    runs = proto.findall(qn("a:r"))
    for extra in runs[1:]:
        proto.remove(extra)
    for tag in ("a:fld", "a:br"):
        for el in proto.findall(qn(tag)):
            proto.remove(el)
    if not runs:
        # paragraph with no run: create one before endParaRPr
        r = proto.makeelement(qn("a:r"), {})
        t = r.makeelement(qn("a:t"), {})
        r.append(t)
        end = proto.find(qn("a:endParaRPr"))
        if end is not None:
            end.addprevious(r)
        else:
            proto.append(r)
    for p in paras:
        tx.remove(p)
    for line in lines or [""]:
        p = copy.deepcopy(proto)
        p.find(qn("a:r")).find(qn("a:t")).text = line
        tx.append(p)


def slide_data(slide):
    """The values in a slide's charts and tables, as short lines of text, so a headline can be written
    from them (test report S-02: a chart slide kept its topic title "Sales by channel")."""
    out = []
    _pics, charts, tables, _other = _visuals(slide)
    for shp in charts:
        try:
            ch = _read_chart(shp)
        except Exception:  # noqa: BLE001  (an unusual chart: no data line, never a failed transform)
            continue
        line = "Chart%s: categories %s" % (" '%s'" % ch["title"] if ch["title"] else "", ", ".join(ch["categories"]))
        for s in ch["series"]:
            line += "; %s: %s" % (s["name"], ", ".join("%g" % v for v in s["values"]))
        out.append(line[:600])
    for shp in tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in list(shp.table.rows)[:12]]
        out.append(("Table: " + "; ".join(rows))[:600])
    return out


def slide_texts(data, indices, whole_slide=False):
    """Title and body text per slide, for the LLM. whole_slide=True (re-templating, where the slide is
    rebuilt from this text) makes body every text box on the slide, so a simplified body replaces all of
    it; otherwise body is the main text box only, the one apply_in_place() rewrites, and the rest is
    read-only other_text."""
    prs = open_deck(data)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        if i not in indices:
            continue
        title = find_title(slide)
        head = title.text_frame.text.strip() if title else ""
        if whole_slide:
            item = {"index": i, "title": head, "body": slide_content(slide, title), "other_text": []}
        else:
            body = find_body(slide, title)
            item = {"index": i, "title": head, "body": _text_of(body) if body else [],
                    "other_text": other_text(slide, title, body)[0]}
        item["data"] = slide_data(slide)        # read-only chart and table values
        item["cover"] = _is_cover(slide)
        out.append(item)
    return out


def apply_in_place(data, changes):
    """changes: {index: {"title": str, "body": [str] or None}}. Returns (bytes, report)."""
    prs = open_deck(data)
    report = []
    for i, slide in enumerate(prs.slides, 1):
        ch = changes.get(i)
        if not ch:
            continue
        notes = []
        title = find_title(slide)
        if ch.get("title") and title is not None:
            old = title.text_frame.text.strip()
            if ch["title"].strip() != old:
                replace_paragraphs(title.text_frame, [ch["title"].strip()])
                notes.append("Headline rewritten")
        elif ch.get("title"):
            notes.append("No title found on this slide; headline not changed")
        if ch.get("body") is not None:
            body = find_body(slide, title)
            if body is not None and ch["body"]:
                if [x.strip() for x in ch["body"]] != _text_of(body):
                    replace_paragraphs(body.text_frame, [x.strip() for x in ch["body"]])
                    notes.append("Main text simplified (%d points)" % len(ch["body"]))
            elif ch["body"]:
                notes.append("No main text box found; body not changed")
        report.append({"index": i, "changes": notes or ["No change needed"], "warnings": []})
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), report


# ------------------------------------------------------------------ re-template
def _read_chart(shp):
    ch = shp.chart
    plot = ch.plots[0]
    name = str(ch.chart_type).upper()
    ctype = "line" if "LINE" in name else "bar" if name.startswith("BAR") else "stacked" if "STACKED" in name else "column"
    cats = [str(c) for c in plot.categories]
    series = [{"name": s.name or "Series %d" % (k + 1), "values": [float(v) if v is not None else 0.0 for v in s.values]}
              for k, s in enumerate(plot.series)]
    title = ""
    if ch.has_title and ch.chart_title.has_text_frame:
        title = ch.chart_title.text_frame.text
    return {"type": ctype, "categories": cats, "series": series, "title": title}


def _add_table(slide, src_shp, x, y, w, h, colors):
    src = src_shp.table
    rows, cols = len(src.rows), len(src.columns)
    gf = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gf.table
    size = 12 if rows <= 8 else 10 if rows <= 14 else 8
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = src.cell(r, c).text
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(size)
                    run.font.bold = r == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else colors["ink"])
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(colors["accent"] if r == 0 else ("FFFFFF" if r % 2 else colors["tile"]))


def _is_cover(slide):
    for ph in slide.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.CENTER_TITLE:
            return True
    return slide.slide_layout.name.strip().lower() in ("title slide", "title", "cover")


def _cover(prs, tpl, head, lines, meta, n, total):
    slide = prs.slides.add_slide(eng.find_layout(prs, tpl, "title"))
    eng.prune_placeholders(slide, TITLE_TYPES + (PP_PLACEHOLDER.SUBTITLE,))
    t = eng.title_shape(slide)
    sub_text = "\n".join(lines[:2]) if lines else meta["date"]
    c = eng.PptxCanvas(slide, tpl)
    if t is not None:
        eng.set_text(t.text_frame, head)
    else:
        c.text(0.6, inches(prs.slide_height) * 0.35, inches(prs.slide_width) - 1.2, 1.5, head, size=36, bold=True)
    sub = next((ph for ph in slide.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE), None)
    if sub is not None:
        eng.set_text(sub.text_frame, sub_text)
    else:
        c.text(0.6, inches(prs.slide_height) * 0.6, inches(prs.slide_width) - 1.2, 0.9, sub_text.split("\n"),
               size=16, color=c.colors["muted"])
    eng.footer(prs, c, 0, total, meta)
    return slide


def _fit(aspect, w, h):
    if w / h > aspect:
        return h * aspect, h
    return w, w / aspect


def retemplate(data, indices, changes, tpl, meta, template_path, template_bytes=None, template_is_potx=False):
    """Rebuild the selected slides on the corporate template (or an uploaded one, if
    template_bytes is given). Returns (bytes, report)."""
    src = open_deck(data)
    prs = eng.open_template_bytes(template_bytes, template_is_potx) if template_bytes is not None else eng.open_template(template_path)
    layout = eng.find_layout(prs, tpl, "content")
    selected = [(i, s) for i, s in enumerate(src.slides, 1) if i in indices]
    total = len(selected)
    report = []
    # small print that repeats across the deck (a template name, a confidentiality mark) is the old
    # template's footer, not something the slide says
    small_by_slide = {i: other_text(s, find_title(s), find_body(s, find_title(s)))[1] for i, s in selected}
    seen = {}
    for lines in small_by_slide.values():
        for x in set(lines):
            seen[x] = seen.get(x, 0) + 1
    repeated = {x for x, k in seen.items() if k > 1}
    for n, (i, s) in enumerate(selected, 1):
        ch = changes.get(i) or {}
        title = find_title(s)
        head = (ch.get("title") or (title.text_frame.text.strip() if title else "") or "Slide %d" % i).strip()
        # the simplified body replaces everything the slide said (it was written from all of it);
        # otherwise carry every text box: KPI tiles and callouts live in boxes of their own
        original = slide_content(s, title)
        lines = list(ch["body"]) if ch.get("body") else original
        # a sentence the new headline was written from must not repeat as a point under it (S-03)
        lines = [x for x in lines if _squash(x) != _squash(head)]
        simplified = bool(ch.get("body")) and [x.strip() for x in ch["body"]] != original
        small_print = [x for x in small_by_slide[i] if x not in repeated]
        pics, charts, tables, other = _visuals(s)
        carried, warnings = [], []

        if _is_cover(s) and not (charts or tables or pics):
            slide = _cover(prs, tpl, head, lines, meta, n, total)
            eng.write_notes(slide, s.notes_slide.notes_text_frame.text if s.has_notes_slide and s.notes_slide.notes_text_frame else "")
            done = ["Moved onto the template's title layout"]
            if title is not None and head != title.text_frame.text.strip():
                done.append("Headline rewritten")
            report.append({"index": i, "changes": done, "warnings": []})
            continue

        slide = eng.add_headline_slide(prs, layout, head, tpl)
        c = eng.PptxCanvas(slide, tpl)
        ax, ay, aw, ah = eng.content_area(prs, slide)
        visuals = charts + tables + pics
        if visuals and lines:
            tw = aw * 0.42
            c.text(ax, ay, tw, ah, lines[:8], size=16, bullet=True)
            vx, vw = ax + tw + 0.3, aw - tw - 0.3
        elif lines:
            c.text(ax, ay, aw, ah, lines[:10], size=18, bullet=True)
            vx, vw = ax, aw
        else:
            vx, vw = ax, aw
        if len(lines) > (8 if visuals else 10):
            warnings.append("Text shortened to fit; %d points moved to the notes" % (len(lines) - (8 if visuals else 10)))

        gap = 0.25
        k = len(visuals)
        side_by_side = not lines and 1 < k <= 3   # no text: use the full width, visuals in a row
        if side_by_side:
            vw_each, vh = (vw - gap * (k - 1)) / k, ah
        else:
            vw_each, vh = vw, ((ah - gap * (k - 1)) / k if k else 0)
        for j, shp in enumerate(visuals):
            bx = vx + j * (vw_each + gap) if side_by_side else vx
            by = ay if side_by_side else ay + j * (vh + gap)
            try:
                if shp in charts:
                    c.chart(bx, by, vw_each, vh, _read_chart(shp))
                    carried.append("Chart rebuilt as a native chart in template colours")
                elif shp in tables:
                    _add_table(slide, shp, bx, by, vw_each, min(vh, 0.4 * len(shp.table.rows) + 0.2), c.colors)
                    carried.append("Table carried over (%d×%d)" % (len(shp.table.rows), len(shp.table.columns)))
                else:
                    aspect = (shp.width / shp.height) if shp.width and shp.height else 16 / 9
                    pw, ph = _fit(aspect, vw_each, vh)
                    slide.shapes.add_picture(io.BytesIO(shp.image.blob), Inches(bx + (vw_each - pw) / 2),
                                             Inches(by + (vh - ph) / 2), Inches(pw), Inches(ph))
                    carried.append("Picture carried over")
            except Exception:  # noqa: BLE001  (unusual charts/images: report instead of failing the deck)
                warnings.append("A %s could not be carried over; copy it from the original" %
                                ("chart" if shp in charts else "table" if shp in tables else "picture"))
        for shp in other:
            if box_text(shp):
                warnings.append("A %s's text was carried over as points, but not its graphic; copy the graphic from the "
                                "original if you need it" % ("group" if shp.shape_type == MSO_SHAPE_TYPE.GROUP else "SmartArt diagram"))
            else:
                warnings.append("A diagram, group or media object was not carried over; copy it from the original")

        eng.footer(prs, c, n, total, meta)
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide and s.notes_slide.notes_text_frame else ""
        extra = lines[(8 if visuals else 10):]
        if extra:
            notes = (notes + "\n\n" + layouts.label({"headline": head, "body": lines}, "Moved from the slide:") + "\n"
                     + "\n".join("- " + x for x in extra)).strip()
        if small_print:
            notes = (notes + "\n\n" + layouts.label({"headline": head, "body": lines}, "Small print on the original slide:") + "\n"
                     + "\n".join("- " + x for x in small_print)).strip()
        eng.write_notes(slide, notes)
        changes_done = ["Moved onto %s" % meta["template_name"]]
        if ch.get("title") and title is not None and ch["title"].strip() != title.text_frame.text.strip():
            changes_done.append("Headline rewritten")
        if simplified:
            changes_done.append("Main text simplified")
        report.append({"index": i, "changes": changes_done + carried, "warnings": warnings})
    cp = prs.core_properties
    cp.title = meta.get("title", "")[:250]
    cp.last_modified_by = "Executive Presentation Studio"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), report
