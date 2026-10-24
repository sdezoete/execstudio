"""python-pptx engine: builds decks on the real corporate template with native charts.

Template files live in data/templates/ (see data/templates.json). Both .pptx and
.potx are accepted. Any sample slides inside the template are removed; only its
masters, layouts, theme colours and fonts are used.
"""
import io
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from . import layouts
from .rtl import is_rtl

EMU_PER_IN = 914400
GEOM = {"rect": MSO_SHAPE.RECTANGLE, "chevron": MSO_SHAPE.CHEVRON, "homePlate": MSO_SHAPE.PENTAGON,
        "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE}
ALIGN = {"l": PP_ALIGN.LEFT, "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
CHART_TYPES = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
               "line": XL_CHART_TYPE.LINE_MARKERS, "stacked": XL_CHART_TYPE.COLUMN_STACKED}
TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)


# ------------------------------------------------------------------ template handling
def open_template(path):
    """Open a registered library template (.pptx or .potx) by file path."""
    if path is None or not Path(path).exists():
        return default_template()
    data = Path(path).read_bytes()
    return open_template_bytes(data, str(path).lower().endswith((".potx", ".potm")))


def open_template_bytes(data, is_potx=False):
    """Open a .pptx or .potx already in memory as an editable Presentation, sample slides
    removed (python-pptx only reads .pptx natively, so a .potx needs its content type swapped
    first). Shared by the library path above and by an ad-hoc uploaded template."""
    if is_potx:
        src = zipfile.ZipFile(io.BytesIO(data))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                blob = src.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    blob = blob.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
                dst.writestr(item, blob)
        data = out.getvalue()
    prs = Presentation(io.BytesIO(data))
    clear_slides(prs)
    return prs


def theme_colors(data):
    """Best-effort accent/ink colours read straight from the uploaded file's own theme1.xml, so
    shapes the studio draws itself (KPI tiles, driver trees, charts...) stay roughly on-brand.
    An uploaded file's theme is not guaranteed to be well-formed, so this never raises -- any
    problem just means fewer colours get overridden from the default palette."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read("ppt/theme/theme1.xml")
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        scheme = etree.fromstring(xml).find(".//a:clrScheme", ns)
        out = {}
        for tag, key in (("dk1", "ink"), ("accent1", "accent")):
            el = scheme.find("a:" + tag, ns) if scheme is not None else None
            if el is None:
                continue
            c = el.find("a:srgbClr", ns)
            if c is not None:
                val = c.get("val")
            else:
                c = el.find("a:sysClr", ns)
                val = c.get("lastClr") if c is not None else None
            if val:
                out[key] = val.upper()
        return out
    except Exception:  # noqa: BLE001  malformed or absent theme: just use the default palette
        return {}


def default_template():
    """python-pptx's built-in template, widened from 4:3 to 16:9."""
    prs = Presentation()
    ratio = 13.333 / 10.0
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    parts = [prs.slide_master] + list(prs.slide_master.slide_layouts)
    for part in parts:
        for xfrm in part.element.iter(qn("a:xfrm")):
            off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
            if off is not None:
                off.set("x", str(int(int(off.get("x")) * ratio)))
            if ext is not None:
                ext.set("cx", str(int(int(ext.get("cx")) * ratio)))
    # Executive style: left-aligned bold titles at the top of content slides.
    lvl1 = prs.slide_master.element.find(".//" + qn("p:titleStyle") + "/" + qn("a:lvl1pPr"))
    if lvl1 is not None:
        lvl1.set("algn", "l")
        d = lvl1.find(qn("a:defRPr"))
        if d is not None:
            d.set("b", "1")
    for lt in prs.slide_layouts:
        for ph in lt.placeholders:
            if ph.placeholder_format.type in (PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.SUBTITLE):
                tx = ph._element.find(qn("p:txBody"))  # noqa: SLF001
                if tx is None:
                    continue
                lst = tx.find(qn("a:lstStyle"))
                if lst is None:
                    lst = etree.Element(qn("a:lstStyle"))
                    tx.insert(1, lst)
                lv = lst.find(qn("a:lvl1pPr"))
                if lv is None:
                    lv = etree.Element(qn("a:lvl1pPr"))
                    lst.insert(0, lv)
                lv.set("algn", "l")
    return prs


def clear_slides(prs):
    lst = prs.slides._sldIdLst  # noqa: SLF001  (python-pptx has no public delete)
    for sld in list(lst):
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)


def _ph_types(layout):
    return {ph.placeholder_format.type for ph in layout.placeholders}


def find_layout(prs, tpl, role):
    """role: 'title' or 'content'. Uses names from templates.json, then sensible defaults."""
    wanted = ((tpl.get("layouts") or {}).get(role) or "").strip().lower()
    all_layouts = list(prs.slide_layouts)
    if wanted:
        for lt in all_layouts:
            if lt.name.strip().lower() == wanted:
                return lt
    if role == "title":
        for lt in all_layouts:
            if PP_PLACEHOLDER.CENTER_TITLE in _ph_types(lt):
                return lt
    else:
        for name in ("title only", "title and content"):
            for lt in all_layouts:
                if lt.name.strip().lower() == name:
                    return lt
        for lt in all_layouts:
            types = _ph_types(lt) - {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
            if types == {PP_PLACEHOLDER.TITLE}:
                return lt
    return all_layouts[0]


def title_shape(slide):
    for ph in slide.placeholders:
        if ph.placeholder_format.type in TITLE_TYPES:
            return ph
    return None


def prune_placeholders(slide, keep):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.type not in keep:
            ph._element.getparent().remove(ph._element)  # noqa: SLF001


def set_text(tf, text, size=None, bold=None, color=None, font=None):
    """Replace a text frame's text, keeping the placeholder's own styling unless overrides are given."""
    tf.text = str(text or "")
    for p in tf.paragraphs:
        if is_rtl(p.text):
            p._p.get_or_add_pPr().set("rtl", "1")  # noqa: SLF001
            p.alignment = PP_ALIGN.RIGHT
        for r in p.runs:
            if size:
                r.font.size = Pt(size)
            if bold is not None:
                r.font.bold = bold
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
            if font:
                r.font.name = font


# ------------------------------------------------------------------ canvas
class PptxCanvas:
    def __init__(self, slide, tpl):
        self.slide = slide
        self.colors = {"accent": tpl.get("accent", "1F47C9"), "ink": tpl.get("ink", "1A2421"),
                       "muted": tpl.get("muted", "5A6763"), "tile": tpl.get("tile", "EDF1EF"),
                       "warn": tpl.get("warn", "D9822B"), "white": "FFFFFF"}
        self.font = tpl.get("font")

    def box(self, x, y, w, h, fill=None, line=None, geom="rect"):
        shp = self.slide.shapes.add_shape(GEOM.get(geom, MSO_SHAPE.RECTANGLE), Inches(x), Inches(y), Inches(w), Inches(h))
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = RGBColor.from_string(line)
            shp.line.width = Pt(0.75)
        else:
            shp.line.fill.background()
        # Drop the theme style reference (shadow/bevel/default text colour); fill, line and text are set explicitly.
        style = shp._element.find(qn("p:style"))  # noqa: SLF001
        if style is not None:
            shp._element.remove(style)  # noqa: SLF001
        return shp

    def text(self, x, y, w, h, lines, size=16, bold=False, color=None, align="l", bullet=False,
             fill=None, line=None, geom="rect", anchor="t", inset=0.08):
        if isinstance(lines, str):
            lines = [lines]
        if fill or line or geom != "rect":
            shp = self.box(x, y, w, h, fill, line, geom)
        else:
            shp = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shp.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = ANCHOR.get(anchor, MSO_ANCHOR.TOP)
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(inset)
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            self._para(p, ln, size, bold, color, align, bullet)
        return shp

    def _para(self, p, text, size, bold, color, align, bullet):
        text = str(text or "")
        rtl = is_rtl(text)
        p.alignment = ALIGN["r"] if rtl and align == "l" else ALIGN.get(align, PP_ALIGN.LEFT)
        p.space_after = Pt(6)
        ppr = p._p.get_or_add_pPr()  # noqa: SLF001
        if rtl:
            ppr.set("rtl", "1")
        if bullet:
            ppr.set("marL", "228600")
            ppr.set("indent", "-228600")
            bu_font = etree.SubElement(ppr, qn("a:buFont"))
            bu_font.set("typeface", "Arial")
            etree.SubElement(ppr, qn("a:buChar")).set("char", "•")
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = RGBColor.from_string(color or self.colors["ink"])
        if self.font:
            f.name = self.font
            if rtl:
                etree.SubElement(r._r.get_or_add_rPr(), qn("a:cs")).set("typeface", self.font)  # noqa: SLF001

    def chart(self, x, y, w, h, spec):
        ctype = CHART_TYPES.get(spec.get("type"), XL_CHART_TYPE.COLUMN_CLUSTERED)
        cd = CategoryChartData()
        cd.categories = [str(c)[:40] for c in spec["categories"]]
        for s in spec["series"]:
            cd.add_series(str(s.get("name") or "Series")[:60], [float(v) for v in s["values"]])
        frame = self.slide.shapes.add_chart(ctype, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = frame.chart
        ch.font.size = Pt(12)
        if self.font:
            ch.font.name = self.font
        multi = len(spec["series"]) > 1
        ch.has_legend = multi
        if multi:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
        if spec.get("title"):
            ch.has_title = True
            ch.chart_title.text_frame.text = str(spec["title"])[:120]
            for p in ch.chart_title.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(13)
                    r.font.bold = True
        else:
            ch.has_title = False
        plot = ch.plots[0]
        if len(spec["categories"]) <= 12:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(11)
            if spec.get("number_format"):
                dl.number_format = spec["number_format"]
                dl.number_format_is_linked = False
        if ctype != XL_CHART_TYPE.LINE_MARKERS:
            plot.gap_width = 80
        palette = [self.colors["accent"], self.colors["ink"], self.colors["muted"], self.colors["warn"]]
        for i, ser in enumerate(plot.series):
            col = RGBColor.from_string(palette[i % len(palette)])
            if ctype == XL_CHART_TYPE.LINE_MARKERS:
                ser.format.line.color.rgb = col
                ser.format.line.width = Pt(2.25)
                ser.smooth = False
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = col
        try:
            va = ch.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = RGBColor.from_string("D9DEDC")
            va.format.line.fill.background()
            va.tick_labels.font.size = Pt(11)
            ch.category_axis.tick_labels.font.size = Pt(11)
        except Exception:  # noqa: BLE001  (axis styling is cosmetic)
            pass
        return True


# ------------------------------------------------------------------ geometry
def inches(v):
    return v / EMU_PER_IN


def content_area(prs, slide):
    """Region below the title placeholder and above the footer lines, in inches."""
    sw, sh = inches(prs.slide_width), inches(prs.slide_height)
    t = title_shape(slide)
    if t is not None and t.width:
        left, width = inches(t.left), inches(t.width)
        top = inches(t.top + t.height) + 0.15
    else:
        left, width, top = 0.6, sw - 1.2, 1.75
    left = max(0.3, left)
    width = min(width, sw - left - 0.3)
    top = min(max(top, 1.2), sh * 0.4)
    bottom = sh - 0.95
    return left, top, width, max(1.5, bottom - top)


def footer(prs, c, idx, total, meta, slide=None):
    sw, sh = inches(prs.slide_width), inches(prs.slide_height)
    c.text(0.5, sh - 0.48, sw * 0.55, 0.35, meta["template_name"], size=9, color=c.colors["muted"])
    if idx:
        c.text(sw - 1.55, sh - 0.48, 1.05, 0.35, "%d / %d" % (idx, total), size=9, color=c.colors["muted"], align="r")
    if slide is None:
        return
    text, warn = layouts.evidence_note(slide)
    if text:
        c.text(0.5, sh - 0.85, sw - 1.0, 0.35, text, size=9, color=c.colors["warn" if warn else "muted"])


def add_headline_slide(prs, layout, headline, tpl):
    """New slide from `layout` with the headline in its title placeholder (or a text box if it has none)."""
    slide = prs.slides.add_slide(layout)
    prune_placeholders(slide, TITLE_TYPES)
    t = title_shape(slide)
    size = tpl.get("title_size", 28)
    if size and len(str(headline).split()) > 12:
        size -= 2
    if t is not None:
        set_text(t.text_frame, headline, size=size)
        t.text_frame.word_wrap = True
    else:
        c = PptxCanvas(slide, tpl)
        c.text(0.55, 0.45, inches(prs.slide_width) - 1.1, 1.15, headline, size=size or 26, bold=True)
    return slide


def write_notes(slide, notes, source="", source_label="Source: "):
    text = (notes or "").strip()
    if source:
        text += ("\n" if text else "") + source_label + source
    if text:
        slide.notes_slide.notes_text_frame.text = text


# ------------------------------------------------------------------ build
def build(title, slides, meta, tpl, template_path=None, template_bytes=None, template_is_potx=False):
    prs = open_template_bytes(template_bytes, template_is_potx) if template_bytes is not None else open_template(template_path)
    lt_title = find_layout(prs, tpl, "title")
    lt_content = find_layout(prs, tpl, "content")
    total = len(slides)
    for i, s in enumerate(slides, 1):
        if s["kind"] == "title":
            slide = prs.slides.add_slide(lt_title)
            keep = TITLE_TYPES + (PP_PLACEHOLDER.SUBTITLE,)
            prune_placeholders(slide, keep)
            t = title_shape(slide)
            if t is not None:
                set_text(t.text_frame, s["headline"])
            sub = next((ph for ph in slide.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE), None)
            subtitle = (s.get("body") or [meta["audience"]])[0] + "\n" + meta["date"]
            c = PptxCanvas(slide, tpl)
            if sub is not None:
                set_text(sub.text_frame, subtitle)
            else:
                c.text(0.6, inches(prs.slide_height) * 0.6, inches(prs.slide_width) - 1.2, 0.9, subtitle.split("\n"),
                       size=16, color=c.colors["muted"])
            if t is None:
                c.text(0.6, inches(prs.slide_height) * 0.35, inches(prs.slide_width) - 1.2, 1.5, s["headline"], size=36, bold=True)
            footer(prs, c, 0, total, meta)
        else:
            slide = add_headline_slide(prs, lt_content, s["headline"], tpl)
            c = PptxCanvas(slide, tpl)
            layouts.draw(c, s, content_area(prs, slide))
            footer(prs, c, i, total, meta, s)
        write_notes(slide, layouts.notes_text(s), (s.get("evidence") or {}).get("source", ""), layouts.label(s, "Source: "))
    cp = prs.core_properties
    cp.title = title[:250]
    cp.author = meta.get("author", "Executive Presentation Studio")
    cp.subject = meta.get("audience", "")
    cp.last_modified_by = "Executive Presentation Studio"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
